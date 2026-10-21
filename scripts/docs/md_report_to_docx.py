#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Konversi laporan markdown -> .docx (heading, tabel, bullet, code block, bold, inline code)."""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

GREEN = RGBColor(0x0B, 0x6B, 0x4F)
NAVY = RGBColor(0x12, 0x3A, 0x5E)
GREY = RGBColor(0x5A, 0x66, 0x72)


def new_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.0); s.bottom_margin = Cm(1.9)
        s.left_margin = Cm(2.0); s.right_margin = Cm(1.8)
    st = doc.styles['Normal']
    st.font.name = 'Segoe UI'; st.font.size = Pt(9.5)
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Segoe UI')
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.08
    for name, size, color, before, after in (
            ('Heading 1', 15, GREEN, 14, 5),
            ('Heading 2', 11.5, NAVY, 10, 4),
            ('Heading 3', 10, GREY, 8, 3)):
        h = doc.styles[name]
        h.font.name = 'Segoe UI'; h.font.size = Pt(size); h.font.bold = True
        h.font.color.rgb = color
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.keep_with_next = True
    return doc


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def add_footer(doc, text):
    p = doc.sections[0].footer.paragraphs[0]
    p.text = ''
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text + '  ·  halaman ')
    r.font.size = Pt(8); r.font.color.rgb = GREY
    for instr, placeholder in (('PAGE', '1'), ('NUMPAGES', '1')):
        if instr == 'NUMPAGES':
            sep = p.add_run(' dari '); sep.font.size = Pt(8); sep.font.color.rgb = GREY
        fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), ' %s ' % instr)
        run = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '16'); rPr.append(sz)
        col = OxmlElement('w:color'); col.set(qn('w:val'), '5A6672'); rPr.append(col)
        run.append(rPr)
        t = OxmlElement('w:t'); t.text = placeholder; run.append(t)
        fld.append(run); p._p.append(fld)


INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`)')


def add_rich(paragraph, text, size=9.5):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = paragraph.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = paragraph.add_run(part[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(8.5)
        else:
            r = paragraph.add_run(part); r.font.size = Pt(size)


def parse_table(lines):
    rows = []
    for ln in lines:
        cells = [c.strip() for c in ln.strip().strip('|').split('|')]
        if all(re.fullmatch(r':?-{2,}:?', c or '-') for c in cells):
            continue
        rows.append(cells)
    return rows


def add_table(doc, rows):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=1, cols=ncol)
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0]
    for i in range(ncol):
        cell = hdr.cells[i]; cell.text = ''
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        txt = rows[0][i] if i < len(rows[0]) else ''
        r = p.add_run(txt); r.bold = True; r.font.size = Pt(8.3)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(cell, '0B6B4F')
    for idx, row in enumerate(rows[1:]):
        cells = t.add_row().cells
        for i in range(ncol):
            cell = cells[i]; cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            add_rich(p, row[i] if i < len(row) else '', size=8.2)
            if idx % 2 == 1:
                shade(cell, 'F2F6F4')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code(doc, block):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    for i, line in enumerate(block):
        r = p.add_run(line)
        r.font.name = 'Consolas'; r.font.size = Pt(8)
        if i < len(block) - 1:
            r.add_break()


def hr(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), 'D5DED9')
    pbdr.append(bottom); pPr.append(pbdr)
    p.paragraph_format.space_after = Pt(4)


def convert(md_path, out_path, footer_text='Laporan Fase 1 · Kelompok 4 · TEK1314'):
    lines = open(md_path, encoding='utf-8').read().splitlines()
    doc = new_doc()
    add_footer(doc, footer_text)

    i = 0
    while i < len(lines):
        line = lines[i]
        raw = line.rstrip()

        if raw.startswith('```'):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                block.append(lines[i]); i += 1
            add_code(doc, block)
            i += 1
            continue

        if raw.strip() == '---':
            hr(doc); i += 1; continue

        m = re.match(r'^(#{1,4})\s+(.*)$', raw)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2).strip(), level=min(level, 3))
            i += 1; continue

        if raw.startswith('|'):
            block = []
            while i < len(lines) and lines[i].startswith('|'):
                block.append(lines[i]); i += 1
            add_table(doc, parse_table(block))
            continue

        m = re.match(r'^(\s*)[-*]\s+(.*)$', raw)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            add_rich(p, m.group(2))
            i += 1; continue

        m = re.match(r'^(\s*)\d+\.\s+(.*)$', raw)
        if m:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(2)
            add_rich(p, m.group(2))
            i += 1; continue

        if not raw.strip():
            i += 1; continue

        p = doc.add_paragraph()
        add_rich(p, raw.strip())
        i += 1

    doc.save(out_path)
    return out_path


if __name__ == '__main__':
    src, dst = sys.argv[1], sys.argv[2]
    print('SAVED', convert(src, dst))
